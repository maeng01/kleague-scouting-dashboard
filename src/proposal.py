"""선수 카드 → 1페이지 스폰서 제안 요약 초안 (python-docx).

버튼 한 번으로 '이 선수를, 이런 근거로, 이 카테고리 브랜드에 제안한다'는
초안을 뽑는다. 최종본이 아니라 담당자가 다듬을 출발점.
"""

from __future__ import annotations

import io
from datetime import date

import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor

from . import config as C
from .agency import AgencyCard
from .brand_fit import CategoryFit, Marketability

_ACCENT = RGBColor(0x2E, 0x7D, 0x32)


def _h(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = _ACCENT
    p.paragraph_format.space_after = Pt(2)
    return p


def build_docx(
    row: pd.Series,
    mkt: Marketability,
    fits: list[CategoryFit],
    card: AgencyCard,
) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(10)

    title = doc.add_paragraph()
    r = title.add_run(f"스폰서십 제안 요약 (초안) — {row['Player']}")
    r.bold = True
    r.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.add_run(
        f"{C.SEASON_LABEL} · {row.get('Squad', '—')} · 스트라이커 · "
        f"{row.get('age_display', '—')} · 신뢰도 {row.get('reliability', '—')} · "
        f"생성일 {date.today().isoformat()}"
    ).italic = True

    _h(doc, "1. 한 줄 요약")
    doc.add_paragraph(
        f"{'★' * card.priority_stars}{'☆' * (5 - card.priority_stars)}  "
        f"추천 카테고리: {', '.join(card.headline_categories) or '재평가 필요'}. "
        f"{card.priority_reason}"
    )

    _h(doc, "2. 경기력 스냅샷 (2026 K리그1 스트라이커 22명 풀 내 percentile · 과정 지표)")
    metrics_line = " · ".join(
        f"{m['label']} {row.get(m['pct']):.0f}p"
        for m in C.SNAPSHOT_METRICS
        if pd.notna(row.get(m["pct"]))
    )
    doc.add_paragraph(metrics_line)
    fin = row.get("finishing_90")
    if pd.notna(fin):
        doc.add_paragraph(f"결정력(득점−xG)/90: {fin:+.2f}  — 결과 지표라 변동성 큼, 참고용")
    doc.add_paragraph(
        f"출전 규모: 90분 환산 {row.get('90s', '—')}경기 (신뢰도 {row.get('reliability', '—')})"
        + ("  ⚠ 표본이 작아 percentile·결정력 해석에 주의" if row.get("small_sample") else "")
    )

    bio_line = []
    if row.get("archetype") and row.get("archetype") != "—":
        bio_line.append(f"아키타입 {row['archetype']}")
    if pd.notna(row.get("height_cm")):
        bio_line.append(f"{int(row['height_cm'])}cm")
    if pd.notna(row.get("preferred_foot")):
        bio_line.append({"right": "오른발", "left": "왼발", "both": "양발"}.get(
            row["preferred_foot"], str(row["preferred_foot"])))
    if pd.notna(row.get("market_value_eur")):
        bio_line.append(f"시장가치 €{int(row['market_value_eur']/1000)}k (Transfermarkt)")
    if pd.notna(row.get("contract_until")):
        bio_line.append(f"계약 만료 {row['contract_until']}")
    if pd.notna(row.get("boot_sponsor")):
        bio_line.append(f"축구화 {row['boot_sponsor']}")
    if bio_line:
        doc.add_paragraph("· 선수 정보: " + " · ".join(bio_line))

    _h(doc, "3. 마케팅 가치")
    doc.add_paragraph(
        f"Marketability {mkt.score}/100 (신뢰도: {mkt.confidence}) — "
        + ", ".join(f"{k} {v:.0f}" for k, v in mkt.components.items())
    )
    if pd.notna(row.get("followers")):
        doc.add_paragraph(
            f"인스타그램 @{row.get('ig_handle')} · 팔로워 약 {int(row['followers']):,}명"
            f" ({row.get('followers_confidence')})"
        )
    if pd.notna(row.get("story_note")) and str(row.get("story_note")).strip():
        doc.add_paragraph(f"스토리: {row['story_note']}")

    _h(doc, "4. 강점")
    for s in card.strengths:
        doc.add_paragraph(s, style="List Bullet")

    _h(doc, "5. 브랜드 카테고리 적합도")
    for f in fits[:3]:
        doc.add_paragraph(
            f"· {f.label}: {f.score}/100"
            + (f"  (근거: {', '.join(f.drivers)})" if f.drivers else ""),
            style="List Bullet",
        )

    _h(doc, "6. 추천 전략")
    for s in card.strategies:
        doc.add_paragraph(s, style="List Bullet")

    _h(doc, "7. 리스크 · 유의사항")
    for rk in card.risks:
        doc.add_paragraph(rk, style="List Bullet")

    disc = doc.add_paragraph()
    dr = disc.add_run("\n" + C.DISCLAIMER)
    dr.font.size = Pt(8)
    dr.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
